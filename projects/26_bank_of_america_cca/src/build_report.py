"""Create the English bank valuation report from the frozen analysis outputs."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'data/processed/python_results.json').read_text());commercial=json.loads((ROOT/'data/processed/python_results_commercial.json').read_text());banks=data['current'];bac=banks[0];reg=data['regression_hc3']
def number(v,kind='m'):
 if v is None:return 'n.a.'
 return f'{v:.1%}' if kind=='%' else f'{v:.2f}x' if kind=='x' else f'${v:.2f}' if kind=='$' else f'{v:,.0f}'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11)
 for prop in st.element.xpath('.//w:pBdr'):prop.getparent().remove(prop)
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True;doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
normal=doc.styles['Normal'];normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
header=sec.header.paragraphs[0];header.text='Bank of America comparable company analysis | 8 October 2026';header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor(0,0,0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;run=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
def p(text):return doc.add_paragraph(text)
def heading(title,new=True):
 para=doc.add_paragraph(title,'Heading 1');para.paragraph_format.page_break_before=new;para.paragraph_format.keep_with_next=True;return para
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,w in enumerate(widths):t.columns[i].width=Inches(w);t._tbl.tblGrid.gridCol_lst[i].set(qn('w:w'),str(int(w*1440)))
 for c,h in zip(t.rows[0].cells,headers):c.text=h
 for row in rows:
  for c,value in zip(t.add_row().cells,row):c.text=str(value)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for i,c in enumerate(row.cells):
   c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   cp.append(borders);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EBED');cp.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.05;para.paragraph_format.left_indent=Pt(0);para.paragraph_format.right_indent=Pt(0);para.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.name='Arial';run.font.size=Pt(10);run.font.bold=ri==0;run.font.hidden=False;run.font.color.rgb=RGBColor(0,0,0)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
doc.add_paragraph('Bank of America Comparable Company Analysis','Title')
p('Frozen valuation and information date 8 October 2026. Prepared for Goran Brkich. All monetary amounts are USD millions unless stated per common share.')
p(f'The selected nine-peer median comparison implies {number(data["weighted_price"],"$")} per Bank of America common share versus a market close of {number(bac["asof_close"],"$")}, a conditional gap of {number(data["weighted_price_change"],"%")}. The gap is modest relative to accounting and peer-selection uncertainty. It does not by itself establish an actionable mispricing.')
p(f'The profitability-conditioned regression implies {number(reg["target_price"],"$")}. Its HC3 95% confidence interval for the conditional mean is {number(reg["target_price_mean_ci95"][0],"$")} to {number(reg["target_price_mean_ci95"][1],"$")}. This interval describes uncertainty in the estimated relationship, not a future trading-price range.')
table(['Measure','Result'],[['Target','Bank of America (BAC)'],['Comparison universe','BAC plus nine listed US banking peers'],['Historical actuals','FY2017–FY2025'],['FY2026 treatment','H1 actuals, June LTM and separate full-year earnings estimate'],['Primary valuation','25% P/E, 25% P/B, 50% P/TBV'],['BAC LTM common EPS',number(bac['eps_ltm'],'$')],['BAC tangible book per share',number(bac['tbv'],'$')]], [3,3.9])
p('Bank deposits and funding expenses are operating features, so equity multiples are the primary convention. A sustainable return on tangible equity is useful for explaining a premium to book, while credit quality, capital and business mix constrain a mechanical comparison.')
heading('Scope and comparison universe')
p('All companies are compared on one frozen common-stock trading date. Financial observations are restricted to disclosures filed by that cutoff. Later retrieval dates document when copies were obtained; they do not authorize later financial information.')
table(['Ticker','Company','Business model'],[[r['ticker'],r['name'],r['group']] for r in banks],[.8,3.2,2.9])
p('The full group spans universal and commercial banking, investment banking, wealth management and consumer cards. Scale alone does not make these companies interchangeable. The alternative commercial cohort excludes Goldman Sachs and Morgan Stanley and recalculates the same native valuation and bivariate regression.')
p('The requested ten-year FY2017–FY2026 horizon contains nine completed annual actuals and a separate FY2026 analyst estimate. Banks have a December year-end, so complete FY2026 actuals did not exist on 8 October. FY2016 is retained only as an opening balance for historical return denominators.')
heading('Current valuation and profitability comparison')
table(['Bank','Close','LTM EPS','P/E','P/B','P/TBV','ROTE','CET1'],[[r['ticker'],number(r['asof_close'],'$'),number(r['eps_ltm'],'$'),number(r['pe'],'x'),number(r['pb'],'x'),number(r['ptbv'],'x'),number(r['rote_ltm'],'%'),number(r['cet1'],'%')] for r in banks],[.8,.85,.9,.8,.8,.9,.9,.95])
p('P/E uses calculated LTM common earnings per diluted average share. P/B and P/TBV use June 2026 point-in-time fundamentals. Negative common earnings do not produce a meaningful P/E. Missing facts remain unavailable rather than entering statistics as zero.')
p('ROTE is a common-earnings return on the two-point average selected tangible common equity at June 2025 and June 2026. It can differ from an issuer return using daily or quarterly average equity and its own adjustment conventions. CET1 uses actual standardized capital ratios, excluding required minima and advanced-approach alternatives.')
p('Goldman Sachs and Morgan Stanley benefit from different capital-markets and wealth exposures; Capital One has a card-heavy credit profile and acquisition effects. These differences should be interpreted alongside any apparent valuation premium.')
heading('Peer median valuation and business interpretation')
table(['Method','Peer median','Implied BAC price','Weight'],[[label,number(data['peer_statistics'][k]['median'],'x'),number(data['implied_prices'][k],'$'),number(w,'%')] for label,k,w in [('P/E','pe',.25),('P/B','pb',.25),('P/TBV','ptbv',.5)]],[1.25,1.7,2.45,1.5])
doc.add_picture(str(ROOT/'docs/valuation_and_returns.png'),width=Inches(6.85))
p('The weighted price is a transparent combination of three separately visible estimates. Median multiples reduce the influence of extremes but cannot remove differences in accounting, loan mix or expected growth. Q25–Q75 implied prices in Excel describe cross-sectional dispersion, not confidence bounds.')
p(f'Excluding GS and MS changes the weighted price to {number(commercial["weighted_price"],"$")}. The alternate result provides a business-mix sensitivity rather than a reason to choose whichever group produces the preferred price.')
heading('Earnings bridge and FY2026 estimate')
p('LTM common income equals FY2025 common income plus H1 2026 common income less H1 2025 common income. The diluted-share average is reconstructed as FY2025 average shares plus half of H1 2026 shares less half of H1 2025 shares. This avoids summing period averages as if they were flows.')
table(['Bank','LTM common income','LTM EPS','FY2026 estimated EPS'],[[r['ticker'],number(r['common_ltm']),number(r['eps_ltm'],'$'),number(data['forecast_fy2026'][i]['eps'],'$')] for i,r in enumerate(banks)],[.9,2.3,1.7,2])
p('The single earnings forecast adds actual H1 2026 to an H2 estimate. Matched H2 2025 revenue, noninterest expense and credit provision are reconstructed from annual and first-half values. Default analyst assumptions are 3% revenue growth, 2% expense growth, no provision growth, a 23% H2 tax rate and a 1% decrease in H2 diluted shares.')
p('Other claims on common income use the H1 2026 run rate. When a standalone provision is not disclosed in the selected tags, the model explicitly labels revenue less operating expense less pretax income as a pretax bridge residual. It is not misrepresented as an issuer-reported provision. The forecast does not invent a December 2026 balance sheet.')
heading('Tangible equity and regulatory capital')
table(['Bank','Common equity','Selected TCE','TBV/share','CET1'],[[r['ticker'],number(r['common']),number(r['tce']),number(r['tbv'],'$'),number(r['cet1'],'%')] for r in banks],[.9,1.65,1.65,1.4,1.3])
p('Common equity uses an explicit issuer disclosure where available, then parent equity less preferred carrying amounts. Preferred par values cannot replace carrying amounts. A rounded reported book value multiplied by end shares is retained only as a disclosed reconstruction.')
p('Selected tangible book uses an issuer reported TBV or TCE measure first. The straight-subtraction version deducts goodwill and other intangibles without issuer deferred-tax adjustments. Both versions are visible in Excel, so differences are not hidden by forcing the values to agree.')
p('PNC reports goodwill and other intangibles together in its reconciliation. The model retains the combined reported amount as raw input and subtracts goodwill once in the financial build. Wells Fargo common equity includes the issuer treatment of preferred capital; the explicit common-equity disclosure takes precedence over a simplified subtraction.')
p('Regulatory CET1 and accounting tangible equity are distinct measures with different deductions and denominators. A higher CET1 ratio alone does not establish a higher equity valuation or eliminate exposure to credit and market risks.')
heading('Credit funding and earnings resilience')
table(['Bank','Gross loans','Deposits','Loans/deposits','Income after 100bp stress'],[[r['ticker'],number(r['loans']),number(r['deposits']),number(r['loans']/r['deposits'] if r['loans'] is not None and r['deposits'] else None,'%'),number(data['credit_stress'][i]['income_100bp'])] for i,r in enumerate(banks)],[.8,1.5,1.5,1.35,1.75])
p('The credit stress deducts incremental annual loan losses of 50, 100 or 200 basis points after the selected tax assumption. It holds loans and tangible equity constant and does not model changes in RWA, deposit demand, lending volumes or capital distributions. These are analytical shocks rather than regulatory stress scenarios or issuer guidance.')
p('Capital One card exposures, commercial loan mix and markets-bank financing exposures mean a uniform basis-point stress has different economic severity. Gross and net loans remain separately labeled. Loan allowances exclude reserves for unfunded commitments unless explicitly part of the selected loan balance.')
p('Loan-to-deposit ratios describe funding structure, while deposit cost, uninsured balances and liquidity composition need separate disclosure. A missing cash classification does not become an invented liquidity coverage or net stable funding ratio. Price returns exclude dividend income, even though cash dividends and repurchases are separately analyzed in the financial schedules.')
heading('Regression uncertainty and influence')
p(f'The bivariate model uses {reg["n"]} non-BAC peers and {reg["df"]} residual degrees of freedom. It regresses peer P/TBV on LTM common income divided by average selected tangible equity. The R-squared is {reg["r2"]:.3f}. A close fitted relationship does not establish causality or guarantee that a peer multiple transfers to BAC.')
table(['Term','OLS estimate','HC3 SE','95% CI'],[[label,f'{reg["coef"][i]:.4f}',f'{reg["se_hc3"][i]:.4f}',f'{reg["ci95"][i][0]:.4f} to {reg["ci95"][i][1]:.4f}'] for i,label in enumerate(['Intercept','ROTE'])],[1.4,1.35,1.35,2.8])
p('HC3 uses the OLS bread matrix and a residual-weighted meat matrix, with each squared residual divided by the square of one minus its leverage. Small-sample inference uses a Student t critical value with n minus two degrees of freedom. The Excel sheet shows scalar calculations for the intercept, slope, residuals, leverage and full covariance.')
multi=data['multivariate_hc3']
if multi:
 table(['Multivariate term','Estimate','HC3 SE'],[[label,f'{multi["coef"][i]:.4f}',f'{multi["se_hc3"][i]:.4f}'] for i,label in enumerate(['Intercept','ROTE','CET1','Log assets'])],[3,1.95,1.95])
 p(f'The multivariate extension has {multi["df"]} residual degrees of freedom. It is an exploratory diagnostic with limited power and potentially unstable coefficients, even with HC3 covariance.')
q=data['bootstrap_price_quantiles'];prices=[x['target_price'] for x in data['leave_one_out']]
p(f'The seeded 5,000-replication pairs bootstrap gives fitted BAC price quantiles of {number(q[0],"$")}, {number(q[1],"$")} and {number(q[2],"$")} at 2.5%, 50% and 97.5%. Omitted-peer estimates range from {number(min(prices),"$")} to {number(max(prices),"$")}. Influence diagnostics require economic review, not automatic deletion of a bank.')
heading('Market return risk and correlations')
table(['Bank','Annual volatility','One-year price return','Daily VaR95','Max drawdown'],[[r['ticker'],number(r['annual_vol'],'%'),number(r['price_return_1y'],'%'),number(r['var95_daily'],'%'),number(r['max_drawdown_1y'],'%')] for r in data['market_risk']],[.85,1.55,1.85,1.25,1.4])
doc.add_picture(str(ROOT/'docs/market_risk.png'),width=Inches(6.85))
p(f'BAC beta against the equal-weight nine-peer bank price benchmark is {data["sector_beta_hc3"]["coef"][1]:.3f}, with HC3 standard error {data["sector_beta_hc3"]["se_hc3"][1]:.3f}. The benchmark excludes BAC and is not the S&P 500. All statistics use the last 252 shared daily observations at the cutoff.')
p('Historical VaR and CVaR are empirical one-day tail measures. Correlations describe shared historical price variation, not interchangeable loan portfolios. Unadjusted returns exclude dividends and cannot be substituted for total shareholder return performance.')
comments={
'BAC':'BAC combines consumer deposits, commercial credit, wealth and capital markets. Its tangible-book premium should be interpreted against sustainable common returns, deposit pricing and credit losses. The model presents both the peer-median and profitability-conditioned estimates because their assumptions differ.',
'JPM':'JPM combines scale, diversified businesses and strong reported returns. Acquired portfolios and capital-market activity complicate a simple valuation premium. The latest filed standardized CET1 disclosure takes precedence over preliminary presentation values and the advanced-approach alternative.',
'C':'Citi has a broad international footprint and an ongoing business reshaping. Restructuring expenses, foreign exchange and the changing capital base affect earnings comparisons. The exact filed CET1 percentage is retained rather than replacing it with a rounded earnings-release headline.',
'WFC':'Wells Fargo has a commercial and consumer lending mix and material historical operating constraints. Explicit reported common equity and tangible equity take precedence over a simplified preferred-stock subtraction. Changes in operating efficiency and credit costs should be evaluated alongside loan growth.',
'GS':'Goldman Sachs has greater dependence on investment banking, trading and market financing than a commercial lender. Its loan disclosure includes financing and held-for-sale categories that differ from retail-bank lending. The commercial cohort provides a transparent sensitivity to this business-model difference.',
'MS':'Morgan Stanley combines institutional securities and wealth management. Its funding, capital intensity and fee revenues differ from a traditional deposit-funded commercial bank. Consolidated loans are selected explicitly rather than substituting only held-for-investment lending.',
'USB':'U.S. Bancorp combines banking with payment services. Selected tangible equity follows its issuer reconciliation, while preferred carrying amounts remain separate from preferred par or liquidation value. Acquisition history and payment-related intangibles affect the relation between ROE and ROTE.',
'PNC':'PNC has material acquisition-related balance-sheet growth. Its FY2026 FirstBank transaction affects goodwill, intangibles and comparable loan balances. The official supplement provides explicit common and tangible equity, including deferred-tax adjustments that differ from straight subtraction.',
'TFC':'Truist history begins with the BB&T predecessor and includes the 2019 merger and subsequent business disposals. Reported historical growth is not organic growth. The model retains transaction effects in GAAP earnings unless a separately documented adjustment is entered.',
'COF':'Capital One has substantial consumer card and auto exposures, together with the Discover acquisition. Acquisition costs, intangible assets and the timing of share issuance affect diluted EPS, tangible book and LTM comparisons. A uniform loan-loss shock does not have the same economic meaning as for commercial peers.'}
for page,subset in [('Company interpretation for BAC JPM C WFC and GS',banks[:5]),('Company interpretation for MS USB PNC TFC and COF',banks[5:])]:
 heading(page)
 for r in subset:
  doc.add_paragraph(r['name'],'Heading 2');p(f'Market {number(r["asof_close"],"$")}; P/TBV {number(r["ptbv"],"x")}; tangible return proxy {number(r["rote_ltm"],"%")}; standardized CET1 {number(r["cet1"],"%")}. '+comments[r['ticker']])
heading('Sources reproduction and interpretation limits')
p('The project includes complete original SEC companyfacts and submissions data, selected complete filing HTML reports, native issuer PDFs, and original daily market-price responses. Original bytes are retained separately from marked review copies. Retrieval dates, accessions, URLs and SHA-256 hashes are recorded.')
p('The Source Crosswalk maps each Excel input cell to its exact period, metric, source label and unit. Exact inline XBRL anchors and verified original table rows provide the strongest links. PDF matches identify the original file and used page; ambiguous or unmatched occurrences remain visible rather than highlighting an arbitrary equal number.')
p('The annotation legend is green for income, blue for balance sheets, yellow for cash distributions and purple for regulatory or note disclosures. Original pagination is retained. HTML originals are clearly distinguished from native PDFs; no reconstructed table replaces a source report.')
p('To reproduce the frozen analysis, extract the project archive, install numpy, scipy and matplotlib, then run python src/analysis.py --cohort 1. Use --cohort 2 for commercial peers. The Python Code sheet contains the full script as plain text for copying into VSCode. The --from-excel option reads saved inputs and assumptions without changing the workbook.')
p('The native Excel model recalculates valuation, earnings forecasts and bivariate HC3. Python Results, resampling outputs and market risk are stored runs. The Audit sheet compares independent baseline calculations; it does not drive a business formula. Verification evidence records input changes, preserved actuals, formula scans, native Notes, row heights and source annotation checks.')
p('The valuation remains conditional on the selected GAAP definitions, book basis, peer group and analyst assumptions. Small peer samples, transaction accounting and cyclical trading income limit precision. Excel desktop behavior should be reviewed when the owner makes later edits; the included checks identify the engines that were available for this delivery.')
doc.core_properties.title='Bank of America Comparable Company Analysis';doc.core_properties.author='Goran Brkich';doc.core_properties.subject='Frozen bank valuation on 8 October 2026'
out=ROOT/'deliverables/Bank_of_America_CCA_Report.docx';doc.save(out);print(out)
