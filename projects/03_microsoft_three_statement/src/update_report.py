"""Append the regression and workbook enhancement section to the Word report."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

PROJECT = Path(__file__).resolve().parents[1]


def append_enhancement_section(doc, project=PROJECT):
    if any(p.text == "13 Historical regression analysis" for p in doc.paragraphs):
        raise ValueError("Enhancement section already exists")
    values = json.loads((project / "evidence/regression/regression_results.json").read_text())
    additions = json.loads((project / "evidence/workbook_enhancements.json").read_text())
    levels, changes = values["levels"], values["first_differences"]
    count = sum(additions["note_counts"].values())

    def p(text):
        return doc.add_paragraph(text)

    def h(text):
        return doc.add_paragraph(text, "Heading 2")

    def table(rows):
        t = doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = [2.65, 2.10, 2.10]
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            b = OxmlElement("w:" + edge)
            for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                b.set(qn("w:" + key), value)
            borders.append(b)
        t._tbl.tblPr.append(borders)
        for i, row in enumerate([["Metric", "Log levels", "Log changes"]] + rows):
            cells = t.rows[0].cells if i == 0 else t.add_row().cells
            pr = cells[0]._tc.getparent().get_or_add_trPr()
            pr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                pr.append(OxmlElement("w:tblHeader"))
            for j, (cell, text) in enumerate(zip(cells, row)):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = text
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    run.bold = i == 0
                if i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:fill"), "DCEAF7")
                    cell._tc.get_or_add_tcPr().append(shd)
        return t

    doc.add_page_break()
    doc.add_paragraph("13 Historical regression analysis", "Heading 1")
    p("The Financial Modeling worksheet estimates the historical relationship between operating cash flow and revenue. It uses the same ten actual years, FY2017 to FY2026, as the operating model. Forecast observations are excluded. The original three-statement assumptions and calculations remain the forecast framework.")
    p("The first specification regresses the natural logarithm of CFO on the natural logarithm of revenue, with an intercept. Its slope is a historical elasticity. The second specification uses annual changes in both logarithms and therefore has nine observations. It tests the association in growth rather than relying on the common trend in levels.")
    interval = lambda v: f"{v['ci95_hc3'][1][0]:.4f} to {v['ci95_hc3'][1][1]:.4f}"
    table([
        ["Observations", str(levels["n"]), str(changes["n"])],
        ["Revenue slope", f"{levels['slope']:.4f}", f"{changes['slope']:.4f}"],
        ["HC3 slope standard error", f"{levels['se_hc3'][1]:.4f}", f"{changes['se_hc3'][1]:.4f}"],
        ["HC3 slope 95% interval", interval(levels), interval(changes)],
        ["HC3 slope p value", f"{levels['p_hc3'][1]:.3g}", f"{changes['p_hc3'][1]:.6f}"],
        ["R squared", f"{levels['r_squared']:.4f}", f"{changes['r_squared']:.4f}"],
        ["Durbin Watson", f"{levels['durbin_watson']:.4f}", f"{changes['durbin_watson']:.4f}"],
        ["Maximum leverage", f"{levels['max_leverage']:.4f}", f"{changes['max_leverage']:.4f}"],
    ])
    p("The levels slope is approximately 1.21, with a narrow in-sample interval. Trending annual series can generate a high fit without a stable predictive relationship. The changes specification has a wider interval and one observation with leverage near 0.80. This makes sensitivity to individual years material.")
    picture = project / "evidence/regression/actual_and_fitted_cfo.png"
    doc.add_picture(str(picture), width=Inches(6.45))
    doc.add_paragraph("Actual CFO and the fitted conditional median from the log levels regression. USD billions. The fitted series is an in-sample diagnostic.", "Caption")

    doc.add_page_break()
    doc.add_paragraph("Regression interpretation and reproduction", "Heading 1")
    h("What HC3 changes")
    p("HC3 uses each squared residual divided by the square of one minus its leverage in the covariance calculation. It changes standard errors, t statistics, p values and confidence intervals. It does not change the OLS coefficients, fitted values or R squared. The worksheet exposes centered regressors, residuals, leverage and observation weights so the calculation can be followed row by row.")
    p("Inference explicitly uses Student t with n minus two degrees of freedom. This is approximate in a sample this small. HC3 addresses heteroskedasticity under the model assumptions. It does not correct serial correlation, shared trends, omitted variables or structural breaks. Durbin Watson and the 4/n leverage reference are descriptive diagnostics here. Neither is treated as evidence that the model is valid for forecasting.")
    h("Using the enhanced workbook")
    p(f"The workbook now has 21 worksheets and {additions['chart_count']} editable charts. New charts cover the four professional views, financial statements, assumptions, segments, operating expenses, working capital, capital assets, financing, taxes, historical data and filing disclosures. Their helper data link to the original cells. Active-case charts and analytical observations update with Assumptions D4.")
    p(f"Expanded native text boxes explain each sheet's calculations, financial interpretation and modeling limits. The workbook contains {count:,} cell Notes on key labels, inputs and calculated values. Hover over an annotated cell in desktop Excel to read its definition and formula references. If hover display is disabled, use Review, Notes, Show or Hide Note. These are Notes rather than discussion threads.")
    h("Python reproduction")
    p(f"On Financial Modeling, copy {additions['python_code_range'].split('!')[1]} into regression_analysis.py in VSCode. The same complete file is included under src in the ZIP. Install NumPy and SciPy, with Matplotlib for the optional figure, then run python regression_analysis.py --out regression_output. It writes coefficient estimates, OLS and HC3 uncertainty, diagnostic CSV files and the comparison figure.")
    p("The code embeds the original historical actuals, so it can run after copying without other project files. For edited observations, export a CSV with year, revenue and cfo columns, then pass --csv input.csv. Changing Excel does not rewrite the embedded Python dataset. The supplied Excel regression calculations and Python matrix implementation were compared independently. The existing saved-workbook statement checks also remain in use.")
    p("Method reference: statsmodels OLSResults.HC3_se documentation at https://www.statsmodels.org/stable/generated/statsmodels.regression.linear_model.OLSResults.HC3_se.html. The Student t critical values are supplied as a visible reference table for degrees of freedom 1 to 30. The information cutoff for the financial data remains 8 October 2026; these presentation and analysis additions are dated 10 October 2026.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    doc = Document(args.input)
    append_enhancement_section(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(str(args.output))


if __name__ == "__main__":
    main()
