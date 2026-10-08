"""Build the explanatory Word report from the supplied financial datasets.

Run after model.py. The document tables and figures use the same processed
data and assumptions as the independent numeric forecast.
"""
from __future__ import annotations

import json
from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT = Path(__file__).resolve().parents[1]
OUT = PROJECT / "deliverables"
FIGURES = PROJECT / "evidence/figures"


def billion(v):
    return f"{v / 1000:,.1f}"


def pct(v):
    return f"{v * 100:.1f}%"


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    data = json.loads((PROJECT / "data/processed/historical_financials.json").read_text())
    forecasts = json.loads((PROJECT / "data/processed/forecast_results.json").read_text())
    assumptions = json.loads((PROJECT / "data/processed/assumptions.json").read_text())
    validation = json.loads((PROJECT / "data/processed/validation_results.json").read_text())
    excel = json.loads((PROJECT / "evidence/excel_verification.json").read_text())
    actual = data["actuals"]["2026"]
    inc, bal, cf = actual["income"], actual["balance"], actual["cashflow"]
    base = forecasts["Base"]
    final = base[-1]
    notes = data["notes"]["2026"]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.7)
    section.bottom_margin = Inches(.7)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    section.footer_distance = Inches(.3)
    styles = doc.styles
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(9)
    styles["Subtitle"].font.size = Pt(12)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(0)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True
    styles["Caption"].font.bold = False
    for style in styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Goran Brkich  |  ")
    r.font.name = "Arial"
    r.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    doc.core_properties.author = "Goran Brkich"
    doc.core_properties.title = "Microsoft Three Statement Financial Model"
    doc.core_properties.subject = "Historical financial statements and integrated forecast"

    def p(text):
        return doc.add_paragraph(text)

    def h(text):
        return doc.add_paragraph(text, "Heading 2")

    def page(title):
        doc.add_page_break()
        doc.add_paragraph(title, "Heading 1")

    def caption(text):
        doc.add_paragraph(text, "Caption")

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths:
            for c, w in zip(t.columns, widths):
                c.width = Inches(w)
        pr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D9D9D9")
            borders.append(el)
        pr.append(borders)
        for row_idx, values in enumerate([headers] + list(rows)):
            cells = t.rows[0].cells if row_idx == 0 else t.add_row().cells
            trpr = cells[0]._tc.getparent().get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            trpr.append(no_split)
            if row_idx == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for col_idx, (cell, value) in enumerate(zip(cells, values)):
                if widths:
                    cell.width = Inches(widths[col_idx])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ["top", "bottom", "left", "right"]:
                    el = OxmlElement("w:" + side)
                    el.set(qn("w:w"), "85")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                tcpr.append(margins)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(1)
                cp.paragraph_format.space_before = Pt(1)
                cp.paragraph_format.line_spacing = 1.0
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                rr = cp.add_run(str(value))
                rr.font.name = "Arial"
                rr.font.size = Pt(10)
                rr.font.color.rgb = RGBColor(0, 0, 0)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DCE6F1" if row_idx == 0 else "F6F8FA" if row_idx % 2 == 0 else "FFFFFF")
                tcpr.append(shade)
                if row_idx == 0:
                    rr.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    doc.add_paragraph("Microsoft Three Statement Financial Model", "Title")
    doc.add_paragraph("Financial analysis and model methodology", "Subtitle")
    p("Goran Brkich\nPortfolio project 03\nInformation cutoff 8 October 2026")
    p("Microsoft's operating growth and its rising investment in infrastructure must be assessed together. This model connects earnings, assets, financing and cash flow so that a change in a forecast driver has a consistent effect on all three financial statements.")
    p(f"The Base case projects revenue of ${billion(final['revenue'])} billion in FY2031. Its operating margin falls from {pct(inc['operating_income']/inc['revenue'])} in FY2026 to {pct(final['operating_income']/final['revenue'])}. Depreciation from new investment explains much of this decline. Free cash flow rises as assumed cash capital expenditure decreases relative to revenue.")
    table(["Metric", "FY2026 actual", "FY2031 Base"], [
        ["Revenue in USD billions", billion(inc["revenue"]), billion(final["revenue"])],
        ["Operating income in USD billions", billion(inc["operating_income"]), billion(final["operating_income"])],
        ["Net income in USD billions", billion(inc["net_income"]), billion(final["net_income"])],
        ["Cash capital expenditure in USD billions", billion(-cf["capex"]), billion(final["cash_capex"])],
        ["Free cash flow in USD billions", billion(cf["cfo"]+cf["capex"]), billion(final["fcf"])],
        ["Diluted earnings per share in USD", f"{inc['eps_diluted']:.2f}", f"{final['eps_diluted']:.2f}"],
    ], [3.45, 1.6, 1.6])
    p("The historical period contains exactly ten fiscal years, FY2017 through FY2026. FY2016 provides opening balances only. The forecast covers FY2027 through FY2031 in Base, Upside and Downside scenarios. Each fiscal year ends on 30 June. Workbook amounts are USD millions, with per-share amounts and ratios labeled separately.")
    p("The workbook includes distinct Equity Research, Investment Banking, Private Equity and FP&A views. The project provides an operating forecast for subsequent valuation work. The investment thesis and valuation will be developed in separate portfolio projects.")

    page("Data sources and historical comparability")
    h("Primary reporting basis")
    p("The income statement, balance sheet and cash flow statement come from Microsoft's annual reports and its FY2026 Form 10-K. Each historical statement uses the latest comparative presentation available by the information cutoff. This avoids combining an original number with a later restatement without recording its origin.")
    p("The long-form historical CSV records the fiscal year, statement, metric, value, unit, source report year, source table index and URL. The source tables and selected SEC XBRL facts are included as extracts. The source manifest records the URLs and hashes of the originally retrieved documents. The complete original HTML documents are obtained with the supplied source-download script when needed.")
    table(["Evidence", "Purpose"], [
        ["Microsoft annual reports", "Audited historical statements and comparative figures"],
        ["FY2026 Form 10-K", "Latest statements, segments, debt and lease disclosures"],
        ["Selected SEC Company Facts", "Independent period and amount controls"],
        ["Source table extracts", "Offline evidence for statement line items"],
    ], [2.45, 4.2])
    h("Accounting changes and segment definitions")
    p("FY2017 contains retrospectively presented revenue and lease accounting figures. An opening statement can come from a different annual report from the following year's income or cash flow statement. Therefore, a historical cash-flow working-capital line is kept as reported, rather than replaced with a simple difference between two separately presented balance sheets.")
    p("The segment series is not a single unchanged reporting definition over the decade. FY2023 through FY2026 use the current definition available in the later reports. Earlier years retain their latest available historical presentation. The workbook marks FY2023 segment growth unavailable across the definition change. Company-wide totals remain useful, but a long segment growth comparison requires attention to the source presentation.")
    h("Disclosures and precision")
    p("Some note figures are disclosed in billions or rounded millions. Depreciation and intangible amortization are kept distinct from the cash-flow line labeled depreciation, amortization and other. The difference is a disclosed presentation issue, not an amount that should be silently forced to zero. Reported data, derived historical ratios and forecast assumptions have separate roles in the model.")

    page("Ten years of historical performance")
    hist_rows = []
    for year in data["historical_years"]:
        a = data["actuals"][str(year)]
        i, c = a["income"], a["cashflow"]
        hist_rows.append([year, billion(i["revenue"]), pct(i["operating_income"]/i["revenue"]), billion(-c["capex"]), billion(c["cfo"]+c["capex"])])
    table(["Fiscal year", "Revenue", "Op margin", "Cash capex", "FCF"], hist_rows, [1.05, 1.4, 1.4, 1.4, 1.4])
    caption("Revenue, cash capital expenditure and free cash flow are in USD billions. FCF equals operating cash flow less cash additions to property and equipment.")
    rev_first = data["actuals"]["2017"]["income"]["revenue"]
    cagr = (inc["revenue"]/rev_first)**(1/9)-1
    p(f"Revenue grows at {pct(cagr)} annually from FY2017 to FY2026, measured across nine year-to-year intervals. The latest year shows revenue growth of {pct(inc['revenue']/data['actuals']['2025']['income']['revenue']-1)}. Operating margins expand over the decade, while the recent investment cycle changes the relationship between earnings and cash generation.")
    p(f"FY2026 operating cash flow is ${billion(cf['cfo'])} billion, up from ${billion(data['actuals']['2025']['cashflow']['cfo'])} billion. Cash capital expenditure rises to ${billion(-cf['capex'])} billion, or {pct(-cf['capex']/inc['revenue'])} of revenue. FCF is consequently ${billion(cf['cfo']+cf['capex'])} billion, below FY2025 despite higher earnings.")
    years = data["historical_years"]
    fig, ax = plt.subplots(figsize=(7, 2.5))
    ax.plot(years, [data["actuals"][str(y)]["income"]["revenue"]/1000 for y in years], color="#183A5E", label="Revenue", linewidth=2)
    ax.plot(years, [(data["actuals"][str(y)]["cashflow"]["cfo"]+data["actuals"][str(y)]["cashflow"]["capex"])/1000 for y in years], color="#4488A5", label="Free cash flow", linewidth=2)
    ax.set_ylabel("USD billions")
    ax.set_xticks(years[::2])
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=.2)
    ax.legend(frameon=False, ncol=2, loc="upper left")
    fig.tight_layout()
    chart = FIGURES / "historical_revenue_and_fcf.png"
    fig.savefig(chart, dpi=220)
    plt.close(fig)
    doc.add_picture(str(chart), width=Inches(6.55))

    page("Revenue scenarios and operating expenses")
    p("The forecast grows each of the three reportable segments separately. Consolidated revenue is the sum of Productivity and Business Processes, Intelligent Cloud and More Personal Computing. It is not entered as a separate independent forecast. Intelligent Cloud is the largest growth contributor in the Base assumptions.")
    table(["Revenue growth assumption", "Base", "Upside", "Downside"], [
        ["PBP FY2027 to FY2031", "14% to 10%", "18% to 12%", "8% to 5%"],
        ["IC FY2027 to FY2031", "24% to 16%", "32% to 20%", "14% to 9%"],
        ["MPC FY2027 to FY2031", "1% to 2%", "4% to 3%", "-5% to 1%"],
        ["Cash capex / revenue", "32% to 20%", "34% to 23%", "35% to 25%"],
        ["Legacy PPE remaining life", "8 years", "9 years", "7 years"],
        ["New PPE useful life", "6 years", "6 years", "5 years"],
    ], [2.75, 1.3, 1.3, 1.3])
    p("These are research assumptions dated 8 October 2026. They are editable hypotheses for the model, rather than company guidance or an investment recommendation. The Upside case pairs stronger growth with higher reinvestment. The Downside case combines slower growth, greater relative investment needs and shorter depreciation lives.")
    h("Cost separation and allocation")
    p("Segment cost rates exclude depreciation, intangible amortization and operating lease expense before these items are added from their separate schedules. This structure makes the infrastructure cost burden visible. Research and development also receives its explicit share of these expenses. Stock-based compensation remains embedded in operating costs and is added back once in operating cash flow.")
    table(["Allocation", "PBP", "IC", "MPC", "R and D"], [
        ["PPE depreciation", "5%", "75%", "5%", "15%"],
        ["Intangible amortization", "20%", "5%", "65%", "10%"],
        ["Operating lease expense", "20%", "60%", "10%", "10%"],
    ], [2.7, .99, .99, .99, .99])
    p("Microsoft does not disclose these expense allocations by segment. They are model assumptions, and each allocation sums to 100%. Their purpose is to explain the forecast economics without treating an estimated allocation as a reported segment fact. Segment operating expense allocation uses the FY2026 reported distribution.")

    page("Working capital and cash conversion")
    p("Receivables depend on revenue and receivable days. Inventory and operating payables depend on cost of revenue and their corresponding days. The initial forecast uses FY2026 closing-balance ratios on a 365-day basis. These are annual modeling ratios, not average-balance turnover statistics.")
    table(["Base driver", "Initial value", "Forecast relationship"], [
        ["Receivable days", f"{assumptions['cases']['Base']['dso'][0]:.1f}", "Receivables = revenue multiplied by days / 365"],
        ["Inventory days", f"{assumptions['cases']['Base']['inventory_days'][0]:.1f}", "Inventory = cost of revenue multiplied by days / 365"],
        ["Operating payable days", f"{assumptions['cases']['Base']['dpo'][0]:.1f}", "Operating payables = cost of revenue multiplied by days / 365"],
        ["Current unearned revenue", pct(assumptions['cases']['Base']['deferred_rev_ratio'][0]), "Liability as a share of revenue"],
        ["Accrued compensation", pct(assumptions['cases']['Base']['comp_ratio'][0]), "Liability as a share of revenue"],
    ], [2.0, 1.1, 3.55])
    h("Capital expenditure payables")
    p(f"FY2026 accounts payable includes ${billion(notes['capex_payables'])} billion related to capital expenditure. Removing this amount leaves operating payables of ${billion(bal['payables']-notes['capex_payables'])} billion. The model forecasts these two obligations separately because their cash-flow classification differs.")
    p("The increase in operating payables contributes to operating cash flow. The change in capital expenditure payables adds to the recognized asset investment and is excluded from operating cash flow. Counting the total increase in accounts payable as operating funding would overstate CFO and would count part of the capital funding twice.")
    h("Cash-flow construction")
    p("Forecast CFO starts with net income, adds noncash depreciation, amortization, stock-based compensation and the modeled debt-accounting amortization, and reflects changes in operating assets, liabilities and deferred taxes. An increase in a receivable consumes cash. An increase in an operating liability provides cash.")
    p("Historical cash-flow adjustments remain the amounts published by Microsoft. They can differ from balance-sheet differences because acquisitions, currency, restatements and other noncash changes also affect balance-sheet amounts. The forecast uses explicit schedules for the items it models and assumes no unannounced acquisitions or currency changes.")

    page("Capital expenditure and asset depreciation")
    p("Cash capital expenditure, accrued capital investment and new finance lease assets are separate inputs. All three can increase PPE, but only cash capital expenditure appears as the modeled cash investment outflow in the current year. Capitalized cash interest is already included in the cash capex budget.")
    table(["Base forecast", "FY2027", "FY2028", "FY2029", "FY2030", "FY2031"], [
        ["Cash capex", *[billion(r["cash_capex"]) for r in base]],
        ["PPE additions", *[billion(r["ppe_additions"]) for r in base]],
        ["PPE depreciation", *[billion(r["depreciation"]) for r in base]],
        ["Closing net PPE", *[billion(r["balance"]["ppe"]) for r in base]],
        ["Finance lease liabilities", *[billion(r["fin_total"]) for r in base]],
    ], [2.0, .93, .93, .93, .93, .93])
    caption("Amounts in USD billions. PPE additions include the change in capital expenditure payables and new finance lease assets.")
    h("Annual depreciation cohorts")
    p("Legacy depreciable net PPE uses an assumed remaining life and is capped at its remaining depreciable balance. Land is held constant. Each forecast year's new additions form a separate cohort with its own useful-life driver. New additions receive half a year's depreciation in their commissioning year. Prior cohorts then receive a full-year charge subject to their remaining balance.")
    p("The workbook reconciles opening PPE, additions and depreciation to closing PPE. It also compares the ending balance with land, legacy depreciable PPE and the remaining balance of every forecast cohort. Gross PPE and accumulated depreciation remain visible rather than being replaced by an unexplained net balance.")
    h("Intangible assets and forecast margin")
    p("The existing intangible assets follow the FY2026 disclosed annual amortization schedule, with each charge capped at the remaining asset balance. Goodwill stays constant and no new acquisition goodwill is assumed. The increase in PPE depreciation is the main reason the forecast can generate more revenue while showing a lower operating margin.")
    p("The model does not forecast disposals, impairment events or a construction-in-progress commissioning pipeline. A half-year convention approximates the annual timing of new capital investment. These simplifications are material when interpreting cloud infrastructure economics.")

    page("Debt leases and liquidity funding")
    h("Contractual face value and carrying amount")
    p(f"The FY2026 debt maturity schedule totals ${billion(notes['debt_face'])} billion of face value. Its carrying amount is ${billion(bal['current_debt']+bal['long_debt'])} billion after discounts and issuance costs, hedge adjustments and debt-exchange premiums. The debt schedule carries these accounting adjustments separately from cash repayment.")
    table(["FY2026 debt reconciliation", "USD millions"], [
        ["Debt face value", f"{notes['debt_face']:,.0f}"],
        ["Less discounts and issuance costs", f"({notes['debt_discount']:,.0f})"],
        ["Less hedge adjustments", f"({notes['debt_hedge_adjustment']:,.0f})"],
        ["Less debt exchange premiums", f"({notes['debt_exchange_premium']:,.0f})"],
        ["Debt carrying amount", f"{bal['current_debt']+bal['long_debt']:,.0f}"],
    ], [4.65, 2.0])
    h("Finance and operating leases")
    p("Legacy lease payments come from the FY2026 contractual schedules. Payments include interest, so the model separates their interest and principal components. Finance lease principal is a financing cash outflow. Operating lease expense and cash payments are reconciled through the ROU asset amortization and the liability principal reduction in CFO.")
    p("New lease assets and obligations are modeled together, with assumed terms and a half-year convention for additions. New operating leases use a simplified amortization and principal profile. This is an annual planning approximation rather than a contract-by-contract effective-interest schedule. Current liabilities use the next year's expected payment profile, including explicitly labeled FY2032 tail assumptions.")
    h("Funding order and interest timing")
    p("The model first uses available cash generation, then liquidates short-term investments above a minimum reserve, and finally calculates the new borrowing needed to meet the minimum cash balance. The research policy is $25 billion of cash and $20 billion of short-term investments. It represents a model funding need, not an available or committed credit facility.")
    p("Liquidity borrowing occurs at year-end so interest begins in the following year. Interest on debt uses opening face value, and investment income uses opening liquid assets. This timing keeps the forecast free of an unintended interest-and-cash circular reference. The three saved scenarios require no incremental liquidity borrowing.")

    page("Integrated forecast statements")
    table(["Base forecast", "FY2027", "FY2028", "FY2029", "FY2030", "FY2031"], [
        ["Revenue", *[billion(r["revenue"]) for r in base]],
        ["Operating income", *[billion(r["operating_income"]) for r in base]],
        ["Net income", *[billion(r["net_income"]) for r in base]],
        ["Operating cash flow", *[billion(r["cfo"]) for r in base]],
        ["Investing cash flow", *[billion(r["cfi"]) for r in base]],
        ["Financing cash flow", *[billion(r["cff"]) for r in base]],
        ["Closing cash", *[billion(r["balance"]["cash"]) for r in base]],
        ["Total assets", *[billion(r["balance"]["total_assets"]) for r in base]],
        ["Total liabilities", *[billion(r["balance"]["total_liabilities"]) for r in base]],
        ["Total equity", *[billion(r["balance"]["total_equity"]) for r in base]],
        ["Balance difference", *[f"{(r['balance']['total_assets']-r['balance']['total_le'])/1000:.2f}".replace("-0.00", "0.00") for r in base]],
    ], [2.0, .93, .93, .93, .93, .93])
    caption("Amounts in USD billions. Cash outflows are shown as negative numbers in this report and in parentheses in the workbook.")
    h("Connections between the three forms")
    p("The income statement determines net income. That result starts the indirect cash flow statement and increases retained earnings. Cash flows determine closing cash on the balance sheet. The asset, debt, lease and equity schedules determine the remaining balance-sheet accounts. The balance sheet is checked independently and contains no unexplained asset or equity balancing entry.")
    h("Tax equity and shares")
    p("Income tax expense is based on an editable effective rate, initially 20%. Current income tax payable stays constant, while long-term income tax and deferred tax liabilities use explicit ratios. This gives an integrated tax forecast without asserting a jurisdiction-specific tax payment schedule or a full forecast of deferred tax assets.")
    p("Paid-in capital increases with stock-based compensation and cash share issuance. Buybacks are allocated between paid-in capital and retained earnings. Dividends are assumed declared and paid in the same fiscal year. Diluted weighted-average shares follow an independent net-change assumption; the model does not infer a repurchase share price from cash buybacks.")
    p("FY2026 recognized investment gains are not extrapolated into forecast earnings. Forecast other income reflects modeled investment income and financing costs. This explains why the initial forecast growth in net income can differ from growth in operating income.")

    page("Comparison of the three scenarios")
    end = [forecasts[c][-1] for c in ["Base", "Upside", "Downside"]]
    table(["FY2031 result", "Base", "Upside", "Downside"], [
        ["Revenue", *[billion(r["revenue"]) for r in end]],
        ["Operating margin", *[pct(r["operating_income"]/r["revenue"]) for r in end]],
        ["Net income", *[billion(r["net_income"]) for r in end]],
        ["Free cash flow", *[billion(r["fcf"]) for r in end]],
        ["FCF after finance lease principal", *[billion(r["fcf_after_finance_leases"]) for r in end]],
        ["Net PPE", *[billion(r["balance"]["ppe"]) for r in end]],
        ["Finance lease liabilities", *[billion(r["fin_total"]) for r in end]],
        ["Cash and short-term investments", *[billion(r["balance"]["liquid_assets"]) for r in end]],
        ["Diluted EPS in USD", *[f"{r['eps_diluted']:.2f}" for r in end]],
    ], [2.8, 1.28, 1.28, 1.28])
    caption("Monetary figures in USD billions except EPS. The table is a saved comparison generated from the supplied Python assumptions. The Excel workbook has one active forecast selected in Assumptions D4.")
    p(f"The Upside case produces less FCF than Base in FY2027 despite higher earnings: ${billion(forecasts['Upside'][0]['fcf'])} billion versus ${billion(base[0]['fcf'])} billion. Its larger investment budget initially absorbs the extra operating cash. By FY2031, stronger operating growth produces more free cash flow, even with higher continuing investment.")
    p("The Downside case materially compresses earnings as slower growth and shorter asset lives increase the cost burden. It still generates positive cash flow under the specified working-capital, payout and reinvestment assumptions. This should not be interpreted as proof that every adverse event would leave liquidity intact.")
    p("The accumulation of cash is especially sensitive to the assumed decline in cash capex intensity and the fixed capital-return rules. Acquisitions, larger buybacks, higher dividends or sustained infrastructure spending would reduce those cash balances. The saved FY2031 liquidity figures are conditional model results.")

    page("Interpretation for four professional roles")
    h("Equity Research")
    p("The Equity Research sheet focuses on sales growth, operating margin, earnings per share, cash conversion and capital intensity. The central research question is whether growth in cloud and productivity products creates cash returns after infrastructure spending and lease obligations. Reported earnings and FCF need separate interpretation because investment gains and noncash charges affect them differently.")
    p("The relevant follow-up work is to test revenue and investment assumptions against product demand, competitive conditions and monetization. The author's own variant view will be developed through an interview before a research recommendation or target price is added in the later Equity Research case.")
    h("Investment Banking")
    p("The Investment Banking sheet distinguishes carrying debt, lease obligations, liquidity and net debt under explicitly different definitions. It presents operating income plus disclosed depreciation and amortization as an analytical earnings proxy. Microsoft does not report this figure as EBITDA, and lease treatment must remain consistent when the model later feeds valuation or financing analysis.")
    p("This view supports the quality of the enterprise-to-equity and capital-structure inputs for a future valuation. It does not estimate a transaction price or a financing commitment.")
    h("Private Equity")
    p("The Private Equity sheet evaluates reinvestment needs, stock-based compensation and cash conversion as operating diligence questions. It compares cash after finance lease principal with accounting earnings and shows how capital intensity affects distributable cash. Microsoft is assessed as a listed operating company; the view does not assume a leveraged acquisition or private ownership.")
    h("FP and A")
    p("The FP&A sheet provides a driver bridge from FY2026 operating income to the selected forecast period. It adds revenue changes by segment and subtracts changes in costs and operating expenses. The bridge updates with the shared forecast drivers, supporting an explanation of the forecast rather than a separate set of contradictory forecasts.")
    p("Its strongest planning use is to identify which assumptions account for a change in profit, cash investment or funding. Budget-versus-actual comparisons and operational budget ownership require a separate case with actual planning data.")

    page("Validation and material modeling limits")
    table(["Validation", "Result"], [
        ["Historical source and reconciliation controls", f"{validation['historical_primary_checks']} passed"],
        ["Forecast relationships across three scenarios", f"{validation['forecast_checks_passed']} passed"],
        ["Recorded Excel versus Python output comparisons", f"{len(excel['checks'])} comparisons"],
        ["Later-year capital expenditure response", "PPE rises and FCF falls when FY2031 capex increases"],
        ["Recorded unselected missing-case input test", excel["unselected_missing_case_test"]],
        ["Recorded formula error scan", "No matches"],
    ], [4.25, 2.4])
    p("The Python forecast was rerun after recovery of the saved files. All 90 forecast relationships passed again. Their maximum absolute difference is below one millionth of a USD million. The original spreadsheet verification records comparisons for all three cases and a later-year input-change test. These records are retained in the evidence directory.")
    p("The supplied verifier reads the saved Excel values and compares statement lines with the Python model. It also checks historical figures, formula-error cells, native text boxes and the direction of dependencies. A saved-value check does not itself execute Excel recalculation. After changing a workbook input, recalculate and save in Excel before using the verifier.")
    h("What the forecast simplifies")
    p("The largest model risks are the future revenue trajectory, the timing of investment and asset commissioning, useful lives, lease additions and segment expense allocations. New lease amortization is an annual approximation. Debt coupons and investment yields are explicit blended assumptions. The forecast omits unannounced acquisitions, disposal and impairment events, currency changes and changes in AOCI.")
    p("Tax liabilities use planning ratios rather than a detailed legal-entity or jurisdictional tax forecast. Diluted shares are an independent assumption. Repurchase cash does not automatically determine the share-count change. No covenant package, committed borrowing facility, credit rating, market valuation or share-price forecast is established by these accounting calculations.")
    p("Excel assumptions and Python assumptions are separately editable. A change made only in Excel is not automatically written to the Python JSON. The project includes an assumption export utility so that a saved Excel scenario can be checked against the same drivers in Python.")

    page("Using the project and locating its sources")
    h("Excel workflow")
    p("Open the workbook, begin with Overview and review the linked statements. On Assumptions, set D4 to 1 for Base, 2 for Upside or 3 for Downside. Forecast drivers are in columns O through S. Each active driver is followed by the three editable scenario rows. Blue numeric entries are inputs. Green entries in the working schedules link to another worksheet. Black entries are local formulas or summary results.")
    p("Review Segments, Working Capital, PPE and Intangibles, Debt and Leases, and Tax and Equity to trace the output. The Audit sheet independently reviews the model and does not determine business results. Native editable text boxes explain calculations next to the related schedules. The four role sheets interpret the same forecast.")
    h("Python and portability")
    p("The forecast runs offline on the supplied historical data. Run python src/model.py from the project directory. The --case option changes the scenario printed in the terminal; all three cases are recalculated and saved on every run. Edit data/processed/assumptions.json to change Python assumptions. Follow README.md for saved-workbook verification, report reproduction and source retrieval.")
    p("The archive includes the workbook, Word report, source extracts, historical and forecast datasets, code, dependencies and verification evidence. Unzip it to retain the relative directory structure. The numerical model uses the Python standard library; source normalization additionally uses lxml and report generation uses python-docx and matplotlib.")
    h("Primary references")
    p("Microsoft Corporation. Form 10-K for the year ended 30 June 2026, filed 29 July 2026. Item 8 contains the three financial statements and the relevant property, lease, debt, intangible and segment disclosures. The exact source table indices are recorded in the datasets.")
    p("https://www.sec.gov/Archives/edgar/data/789019/000119312526323660/msft-20260630.htm")
    p("Microsoft Corporation. Annual reports for FY2016 through FY2025. FY2016 is used for the opening balance context. Source URLs and the report chosen for each statement are recorded in historical_financials.csv and source_manifest.json.")
    p("https://www.microsoft.com/en-us/Investor/annual-reports.aspx")
    p("U.S. Securities and Exchange Commission. Microsoft Company Facts, CIK 0000789019. Selected 10-K observations are filtered by fiscal period and the fixed information cutoff.")
    p("https://data.sec.gov/api/xbrl/companyfacts/CIK0000789019.json")
    result = OUT / "Microsoft_Three_Statement_Model_Report.docx"
    doc.save(result)
    print(result)


if __name__ == "__main__":
    build()
