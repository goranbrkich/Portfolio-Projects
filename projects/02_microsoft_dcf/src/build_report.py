"""Create the portfolio research report from independently computed results."""
from pathlib import Path
import json
import math
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
D=json.loads((ROOT/'data/processed/model_inputs.json').read_text())
R=json.loads((ROOT/'data/processed/reference_results.json').read_text())
B=json.loads((ROOT/'data/processed/beta_diagnostics.json').read_text())
base=R['cases']['Base'];rev=R['reverse']['Base'];market=D['market']['price']
figdir=ROOT/'docs/charts';figdir.mkdir(exist_ok=True,parents=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','savefig.dpi':170})
navy='#20364F';teal='#3D7FA6';brown='#AA7655'

def charts():
    fig,ax=plt.subplots(figsize=(7,3.1));names=['Downside','Base','Upside'];values=[R['cases'][x]['fair_price'] for x in names]
    ax.bar(names,values,color=[brown,navy,teal],width=.55);ax.axhline(market,color='#6A6A6A',ls='--',lw=1.5,label=f'8 October close ${market:.2f}')
    for i,x in enumerate(values):ax.text(i,x+10,f'${x:.2f}',ha='center')
    ax.set_ylabel('USD per share');ax.set_ylim(0,600);ax.legend(loc='upper left',frameon=False);ax.set_title('DCF outcomes under the three fiscal scenarios');fig.tight_layout();fig.savefig(figdir/'case_values.png');plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(7,3.2));years=np.arange(2027,2032)
    for name,color in [('Base',navy),('Upside',teal),('Downside',brown)]:
        f=R['cases'][name]['forecast'];axes[0].plot(years,np.array(f['revenue'])/1000,label=name,color=color,marker='o');axes[1].plot(years,np.array(f['fcff'])/1000,color=color,marker='o')
    axes[0].set_title('Revenue');axes[1].set_title('Unlevered free cash flow')
    for a in axes:a.set_ylabel('USD billions');a.set_xticks(years);a.tick_params(axis='x',labelsize=8)
    axes[0].legend(frameon=False,fontsize=8);fig.tight_layout();fig.savefig(figdir/'forecast_paths.png');plt.close(fig)
    fig,ax=plt.subplots(figsize=(7,3));ax.hist(R['monte_carlo']['prices'],bins=28,color=teal,edgecolor='white');ax.axvline(market,color=brown,lw=2,label=f'Observed close ${market:.2f}');ax.set_xlabel('Conditional value USD per share');ax.set_ylabel('Draw count');ax.legend(frameon=False);ax.set_title('Fixed seed triangular simulation with 1000 draws');fig.tight_layout();fig.savefig(figdir/'simulation.png');plt.close(fig)

def main():
    charts();doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7);section.left_margin=section.right_margin=Inches(.75)
    section.header_distance=section.footer_distance=Inches(.3)
    for key in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        s=doc.styles[key];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(8)
    doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(17);doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
    doc.styles['Caption'].font.size=Pt(8.5);doc.styles['Caption'].font.italic=True
    for s in doc.styles:
        for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;run=foot.add_run('Goran Brkich | Portfolio project 02 | ');run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    doc.core_properties.author='Goran Brkich';doc.core_properties.title='Microsoft DCF and Reverse DCF';doc.core_properties.subject='Five year valuation and market implied expectations as of 8 October 2026'
    def p(text):return doc.add_paragraph(text)
    def h(text):return doc.add_paragraph(text,'Heading 2')
    def page(text):doc.add_page_break();doc.add_paragraph(text,'Heading 1')
    def image(name,width=6.8):doc.add_picture(str(figdir/name),width=Inches(width))
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for i,values in enumerate([headers]+list(rows)):
            cells=t.rows[0].cells if i==0 else t.add_row().cells
            trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if i==0:trpr.append(OxmlElement('w:tblHeader'))
            for j,(cell,text) in enumerate(zip(cells,values)):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F7' if i==0 else ('F6F8FA' if i%2 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shade)
                cp=cell.paragraphs[0];cp.paragraph_format.space_after=Pt(4);cp.paragraph_format.space_before=Pt(4);r=cp.add_run(str(text));r.font.size=Pt(9);r.bold=i==0
        p('');return t
    doc.add_paragraph('Microsoft DCF and Reverse DCF','Title');p('Portfolio project 02 | Goran Brkich');p('Valuation and information cutoff 8 October 2026\nFive explicit fiscal years FY2027 to FY2031\nReport prepared 12 October 2026')
    h('Valuation finding')
    p(f"The selected base assumptions produce a value of ${base['fair_price']:.2f} per share against the observed close of ${market:.2f}. The conditional valuation is {abs(base['upside']):.1%} below the close. Upside and downside fiscal cases produce ${R['cases']['Upside']['fair_price']:.2f} and ${R['cases']['Downside']['fair_price']:.2f}. These differences measure the effects of the visible forecasts and valuation conventions.")
    image('case_values.png');doc.add_paragraph('Source Internal DCF calculations and preserved 8 October market observation','Caption')
    p(f"The base case discount rate is {base['wacc']:.2%}, stable growth 3.0%, stable operating margin 40.0% and stable ROIC 20.0%. The stable period supplies {base['terminal_share']:.1%} of enterprise value. The result is therefore particularly sensitive to long-run cash conversion and the required return.")
    p('The reverse analysis explains the market price through one criterion at a time. Revenue CAGR and lower WACC can solve it mathematically, while capex, working capital, tax and lease acquisition tests produce values outside reasonable operating bounds. No positive finite stable ROIC can explain the price alone under the base margins and growth.')

    page('Operating forecast and investment')
    p('The annual forecast retains the project 03 fiscal scenario calibration for segment growth, operating margin, operating tax, D&A, cash capital spending, capitalized interest, finance lease assets and operating working capital. The new DCF adds explicit discount timing and stable growth, margin and ROIC assumptions.')
    image('forecast_paths.png');doc.add_paragraph('Full fiscal year amounts before the first year remaining cash flow adjustment','Caption')
    f=base['forecast'];table(['USD billions','FY2027','FY2028','FY2029','FY2030','FY2031'],[[label]+[f'{x/1000:,.1f}' for x in f[key]] for label,key in [('Revenue','revenue'),('Operating profit','ebit'),('NOPAT','nopat'),('Economic capex','economic_capex'),('D&A','da'),('Change in NWC','change_nwc'),('Unlevered FCF','fcff')]],[1.7,1.02,1.02,1.02,1.02,1.02])
    h('Economic capital spending')
    p('New finance lease assets are included in reinvestment even when acquisition is noncash. Existing finance lease liabilities are deducted in the enterprise-to-equity bridge. Capitalized financing interest is removed from reported cash capex. Operating leases remain inside operating expenses, and their liabilities are removed from operating NWC without an additional financing deduction.')
    p('SBC remains expensed in EBIT and is never added back to FCFF. The displayed 3.75% compensation diagnostic retains the original forecast calibration. Future buybacks are not an additional source of value; funded repurchases consume equity value.')

    page('Discounting stable value and the equity bridge')
    h('Remaining cash flow and timing')
    p('The valuation date is 8 October 2026. There are 265 days until 30 June 2027. The model scales annual FY2027 FCFF by 265/365 and discounts the remaining cash flow at the midpoint of that interval. Later explicit cash flows use fiscal midyear timing. A visible timing switch allows fiscal year-end discounting. Stable value is always discounted from 30 June 2031.')
    h('Stable investment requirement')
    p('Normalized FY2032 NOPAT equals FY2031 revenue times one plus stable growth, stable operating margin and one minus the stable tax rate. Reinvestment equals NOPAT times growth divided by ROIC. Normalized FCFF equals NOPAT less that reinvestment. Terminal value equals normalized FCFF divided by WACC less growth.')
    p('The primary method requires WACC greater than stable growth and ROIC greater than growth. A simple Gordon method that grows final explicit FCFF is shown as a crosscheck. An editable exit multiple provides a separate terminal convention. Neither replaces the primary reinvestment requirement without an explicit analyst choice.')
    bridge=R['bridge'];table(['Valuation component','USD millions'],[['Explicit FCFF present value',f"{base['pv_explicit']:,.0f}"],['Stable value present value',f"{base['pv_terminal']:,.0f}"],['Operating enterprise value',f"{base['enterprise_value']:,.0f}"],['Excess liquid assets',f"{bridge['excess_cash']:,.0f}"],['Equity investment value proxy',f"{bridge['investments']:,.0f}"],['Debt fair value deduction','(36,500)'],['Finance lease liability deduction','(66,594)'],['Equity value',f"{base['equity_value']:,.0f}"],['Million shares',f"{bridge['shares']:,.3f}"],['USD per share',f"{base['fair_price']:,.2f}"]],[4.8,2.0])
    h('Balance date and ownership conventions')
    p('Cash, investments, debt and leases are June disclosure proxies. Debt uses disclosed fair value, equity investments use carrying value times an editable factor, and excess liquid assets are measured after a USD 25 billion operating reserve. The share denominator combines 7,425.545491 million July point shares with a 24 million incremental weighted-share proxy. This is an explicit estimate of existing award overhang rather than an exact treasury-stock option calculation.')

    page('Market implied operating expectations')
    p('Each reverse sheet solves the same market enterprise value with the same bridge. It changes one criterion and holds the other assumptions fixed. The solutions are separate diagnostic questions; combining them would create a different model.')
    table(['Criterion','Base required result','Economic interpretation'],[
        ['Uniform five year revenue CAGR',f"{rev['growth']:.2%}",'Within displayed 0% to 35% analyst band'],
        ['Parallel operating margin shift',f"+{rev['margin_shift']:.2%}",'Explicit and stable margins all rise by this amount'],
        ['Stable margin alone',f"{rev['terminal_margin_only']:.2%}",'Requires unusually high enduring operating margin'],
        ['Explicit margin shift alone',f"+{rev['forecast_margin_only_shift']:.2%}",'Implies operating margins above 100%'],
        ['Cash capex multiplier',f"{rev['capex_scale']:.2f}x",'Negative spending is rejected'],
        ['Operating NWC ratio shift',f"{rev['nwc_shift']:.2%}",'Outside +/-10 pp displayed analyst band'],
        ['Receivables only equivalent',f"{rev['dso_change_days']:.0f} days",'Algebraic DSO equivalent outside feasible collection days'],
        ['Finance lease asset multiplier',f"{rev['lease_scale']:.2f}x",'Negative asset acquisition is rejected'],
        ['Operating tax rate shift',f"{rev['tax_shift']:.2%}",'Implies negative stable and explicit tax rates']],[2.2,1.3,3.3])
    h('Growth and profitability tradeoffs')
    p('The revenue test replaces the three segment paths with one constant CAGR from FY2026. Selected margin, tax, investment and NWC ratios remain fixed. A high revenue result therefore also scales investment and working capital needs. It does not assume that higher revenue arrives without cash investment.')
    p('The parallel margin test adds the same percentage-point amount to each explicit margin and the stable margin. Separate tests isolate explicit profitability and stable profitability. Required margins must be checked against operating bounds and the costs of achieving those improvements.')
    h('Capital intensity and tax tests')
    p('Capex, finance lease additions and NWC are linear cash-flow sensitivities with D&A and stable ROIC held fixed. They are not new asset capacity plans. The tax test shifts operating taxes while holding the WACC debt shield fixed. Negative solutions identify which isolated interpretation cannot explain the price.')

    page('Market implied valuation and stable economics')
    table(['Criterion','Required result','Interpretation'],[
        ['WACC',f"{rev['wacc']:.3%}",f"Compare with calculated {base['wacc']:.3%}"],
        ['Stable growth',f"{rev['terminal_growth']:.3%}",'Outside displayed 0% to 5% long-run analyst band'],
        ['Stable ROIC','No finite positive value','Required cash exceeds stable NOPAT ceiling'],
        ['Price as ROIC tends to infinity',f"${rev['roic_limit_price']:.2f}",f"Still below observed ${market:.2f}"],
        ['Normalized FY2032 FCFF',f"${rev['required_terminal_fcff']/1000:,.1f}bn",f"{rev['required_terminal_fcff_margin']:.2%} of stable revenue"],
        ['Exit EV to FY2031 EBITDA',f"{rev['required_exit_multiple']:.2f}x",'Requires independently supported peer and rate context'],
        ['Shares with equity held constant',f"{rev['required_shares']:,.1f}m",'Cost-free share removal is not an economic buyback']],[2.55,1.45,2.8])
    h('ROIC ceiling')
    p('At a fixed stable margin and growth rate, higher ROIC reduces investment needed for growth. As ROIC tends to infinity, reinvestment tends to zero and FCFF approaches NOPAT. The market target under the base assumptions requires more cash than that ceiling. A negative algebraic ROIC is therefore not reported as a valid business solution.')
    h('Native sensitivity and case comparison')
    p('The DCF sheet contains native two-input Excel What-If tables. The WACC and stable-growth inputs are the actual cells used to price the primary DCF, on the same worksheet. The live case grid temporarily tests the same scenario selector and growth input. The diagonal gives each case at its own stable growth without duplicated forecasts or fixed case snapshots.')
    h('Ownership and exit multiple')
    p('Ownership compares point shares, the point count plus the award proxy and historical diluted EPS shares. A reverse share count is shown only as algebra; the funded buyback example deducts its cash cost. The exit multiple crosscheck uses an explicitly editable illustrative 12x assumption, with no invented comparable-company sample.')

    page('Beta inference and conditional simulation')
    h('Five year market model')
    p(f"The regression has {B['n']:,} daily return observations, ending 8 October 2026, for Microsoft against the SPY S&P 500 proxy. OLS beta is {B['beta']:.4f} and R squared is {B['r_squared']:.3f}. HC3 beta standard error is {B['hc3_se_beta']:.4f}; the approximate 95% Student t confidence interval is {B['hc3_beta_lower']:.4f} to {B['hc3_beta_upper']:.4f}.")
    p(f"HC3 uses squared residuals divided by one minus leverage squared in the covariance sandwich. It changes inference, not the fitted beta. Newey-West lag five beta standard error is {B['hac5_se_beta']:.4f}, compared with conventional OLS {B['ols_se_beta']:.4f}. The workbook also shows a two-year daily beta and optional Blume adjustment. The intercept is a daily total-return market-model alpha, not a historical excess-return CAPM alpha.")
    h('Discount rate sources')
    p('The model uses the 8 October 10-year Treasury yield of 5.22%, explicitly dated 1 October Damodaran ERP of 4.20% and January AAA synthetic spread of 0.40%. Equity weights use cutoff price times point shares. Debt and finance lease liabilities are disclosure proxies. The debt shield uses an explicit 21% marginal tax assumption.')
    image('simulation.png');doc.add_paragraph('Independent triangular priors and fixed uniforms with seed 20261008','Caption')
    mc=R['monte_carlo'];p(f"Under these subjective priors, the conditional median value is ${mc['percentiles']['0.5']:.2f}, with 5th and 95th percentiles ${mc['percentiles']['0.05']:.2f} and ${mc['percentiles']['0.95']:.2f}. The simulation uses 1,000 identical frozen draws in both engines. Those percentiles describe the selected model assumptions, not empirically estimated market outcome probabilities.")
    p('The simulation replaces segment growth with uniform CAGR, changes operating and stable margins together, scales cash capex, and varies WACC and stable growth. D&A, other ratios, bridge and stable ROIC remain fixed. Independent triangular ranges omit correlated shocks, model error and tail events; the mode preserves final revenue but changes the intervening path.')

    page('Raw evidence and reproducible calculations')
    h('Sources and marked materials')
    p('The project includes full original annual report PDFs FY2017 to FY2026 and separate marked copies. The 402-record source crosswalk identifies selected input, fiscal year, report year, native PDF page, raw units, conversion and Excel destinations. All 344 reported records are highlighted; 51 unavailable note disclosures remain n.a. Seven normalized absence records are distinguished from explicitly reported zero values.')
    p('Annotations preserve native PDF text and page counts. Clicking a colored highlight shows its source ID and current DCF cell references. Bookmarks and margin stripes identify used pages. FY2017 is retained as an original archive report, while selected FY2017 observations come from later comparative reports. Income, balance, cash-flow, notes and segment inputs use the documented color legend.')
    p('Raw market materials include original HTML, JSON, CSV and XLSX responses, dated metadata and SHA-256 hashes. The Treasury and Yahoo datasets can contain later information in their full response; only observations at or before 8 October enter the processed inputs. The ERP monthly archive has an inconsistent final date serial, so the dated author homepage determines the period selection. Market_Source_Extracts.pdf is clearly identified as a generated review extract.')
    h('Two implementations')
    p('The workbook performs the live DCF, robust beta inference, reverse criteria and simulation with ordinary formulas and native data tables. Python Code contains the full plain-text Python program, preserving indentation and blank lines. The same program is available as src/valuation.py and uses NumPy and SciPy without requiring embedded Python in Excel.')
    p('For VSCode, keep the complete archive folder structure, install requirements.txt and run python src/valuation.py. To replay edited Excel drivers, run python src/valuation.py --excel deliverables/Microsoft_DCF_ReverseDCF_Model.xlsx. Mapped numeric assumptions are read directly from the XLSX; Python does not use formula caches to obtain calculated values.')
    h('Verification and use')
    p('The Audit sheet compares independently computed scenario values, OLS beta, HC3 and HAC inference, reverse valuation residuals and deterministic simulation results. Enterprise-to-equity and segment-sum identities are also checked. Frozen baseline comparisons intentionally change after editing inputs; rerun Python to verify the edited model.')
    p('Table rows are consistently 24 pt. Expanded observations use merged blocks of those rows. Native hover Notes explain inputs and formulas; charts use linked ranges. The workbook avoids TEXT-based narrative formulas. Calculation uses the XLSX international formula syntax. After editing in Excel, enable Automatic calculation including data tables and recalculate.')
    h('Primary references')
    p('Microsoft annual reports and SEC filing information in the per-cell crosswalk. Issuer price lookup microsoft.gcs-web.com. US Treasury daily par yield curve home.treasury.gov. Damodaran dated ERP and rating spread tables pages.stern.nyu.edu/~adamodar. HC3 method documentation statsmodels.org. Yahoo price responses are retained as the original downloaded market dataset.')
    output=ROOT/'deliverables/Microsoft_DCF_ReverseDCF_Report.docx';doc.save(output);print(output)

if __name__=='__main__':main()
